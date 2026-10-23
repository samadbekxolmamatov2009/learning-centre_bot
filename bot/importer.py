"""Excel (.xlsx) va Word (.docx) fayldan o'quvchilar ro'yxatini o'qish."""
import io
import re

from openpyxl import load_workbook

MAX_ROWS = 3000

HEADERS = {
    "last": {"familiya", "familiyasi", "фамилия", "surname", "lastname", "last", "f"},
    "first": {"ism", "ismi", "имя", "name", "firstname", "first"},
    "grade": {"sinf", "sinfi", "класс", "class", "grade"},
    "phone": {"telefon", "telefonraqam", "telefonraqami", "tel", "telraqam", "raqam", "телефон", "phone", "tell"},
    "full": {"fio", "fish", "ismfamiliya", "familiyaism", "familiyaismi", "ismfamiliyasi", "oquvchi",
             "oquvchining", "ф.и.о", "фио", "student", "fullname", "ismsharif"},
}


def _clean(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def _key(v: str) -> str:
    return re.sub(r"[^a-zа-яё]", "", v.lower().replace("ʻ", "").replace("'", "").replace("’", "").replace("`", ""))


def _digits(v: str) -> str:
    return "".join(ch for ch in v if ch.isdigit())


def _looks_phone(v: str) -> bool:
    return len(_digits(v)) >= 9 and not re.search(r"[a-zа-я]", v.lower())


def _header_map(cells):
    m = {}
    for i, c in enumerate(cells):
        k = _key(_clean(c))
        for field, names in HEADERS.items():
            if k in {_key(n) for n in names} and field not in m:
                m[field] = i
    return m


def _from_headerless(cells):
    """Sarlavhasiz qator: [№], Familiya Ism | Familiya, Ism, [sinf], [telefon] ..."""
    cells = [_clean(c) for c in cells]
    cells = [c for c in cells if c]
    if cells and re.fullmatch(r"\d{1,4}[.)]?", cells[0]) and len(cells) > 1:
        cells = cells[1:]
    if cells:
        cells[0] = re.sub(r"^\d{1,4}[.)]\s+", "", cells[0])  # "1. Valiyev Ali"
    if not cells:
        return None
    if len(cells[0].split()) >= 2:
        name, rest = cells[0], cells[1:]
    elif len(cells) >= 2:
        name, rest = cells[0] + " " + cells[1], cells[2:]
    else:
        return {"full": cells[0], "grade": "", "phone": ""}
    phone = next((c for c in rest if _looks_phone(c)), "")
    grade = next((c for c in rest if c != phone and not _looks_phone(c)), "")
    return {"full": name, "grade": grade, "phone": phone}


def _split_full(full: str):
    parts = full.split()
    if len(parts) < 2:
        return None
    return parts[0], " ".join(parts[1:])


def _rows_to_students(rows, first_line=1):
    """rows: list of lists of cell values. Sarlavhani o'zi topadi."""
    students, errors = [], []
    start, hm = 0, {}
    for i, r in enumerate(rows[:5]):
        hm = _header_map(r)
        if hm and (("last" in hm and "first" in hm) or "full" in hm):
            start = i + 1
            break
        hm = {}
    for i, r in enumerate(rows[start:], start=start):
        line = first_line + i
        if not any(_clean(c) for c in r):
            continue
        if hm:
            get = lambda f: _clean(r[hm[f]]) if f in hm and hm[f] < len(r) else ""
            last, first = get("last"), get("first")
            if get("full") and not (last and first):
                sp = _split_full(get("full"))
                last, first = sp if sp else ("", "")
            grade, phone = get("grade"), get("phone")
        else:
            d = _from_headerless(r)
            if not d:
                continue
            sp = _split_full(d["full"])
            last, first = sp if sp else ("", "")
            grade, phone = d["grade"], d["phone"]
        if not last or not first:
            errors.append((line, "ism yoki familiya topilmadi"))
            continue
        students.append({"line": line, "last": last, "first": first, "grade": grade, "phone": phone})
    return students, errors


def parse_xlsx(data: bytes):
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    return _rows_to_students(rows)


def parse_docx(data: bytes):
    from docx import Document
    doc = Document(io.BytesIO(data))
    students, errors = [], []
    if doc.tables:
        offset = 1
        for t in doc.tables:
            rows = [[c.text for c in row.cells] for row in t.rows]
            s, e = _rows_to_students(rows, first_line=offset)
            students += s
            errors += e
            offset += len(rows)
        return students, errors
    rows = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        parts = re.split(r"[;\t|,]", text)
        if len(parts) == 1:  # "Valiyev Ali 7 901234567" - bo'shliq bilan yozilgan
            toks = text.split()
            toks[0] = re.sub(r"^\d{1,4}[.)]$", "", toks[0]) or ""
            toks = [t for t in toks if t]
            parts = [" ".join(toks[:2])] + toks[2:]
        rows.append(parts)
    return _rows_to_students(rows)


def parse_file(filename: str, data: bytes):
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        students, errors = parse_xlsx(data)
    elif name.endswith(".docx"):
        students, errors = parse_docx(data)
    else:
        raise ValueError("Faqat .xlsx (Excel) yoki .docx (Word) fayl qabul qilinadi.")
    if len(students) + len(errors) > MAX_ROWS:
        raise ValueError(f"Fayl juda katta (ko'pi bilan {MAX_ROWS} qator).")
    return students, errors
