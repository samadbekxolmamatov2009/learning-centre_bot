import asyncio
import calendar
import io
import logging
from datetime import datetime, timedelta, timezone

from aiogram import Router, F, Bot
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import Message, CallbackQuery, BufferedInputFile, InlineKeyboardMarkup, InlineKeyboardButton
from aiogram.utils.keyboard import InlineKeyboardBuilder

from openpyxl import Workbook
from openpyxl.styles import Font

from . import db as D
from .config import TZ_OFFSET, REMINDER_HOUR
from .handlers import (Role, STAFF, ANYONE, DIRECTOR, prev_months, with_phone, group_picker, confirm_kb, money, months_markup, debtors_text,
                       can_access_group, normalize_phone)

router = Router()
log = logging.getLogger(__name__)


# ---------------- O'quvchini boshqa guruhga ko'chirish ----------------
@router.message(F.text == "🔀 O'quvchini ko'chirish", STAFF)
async def mv_start(m: Message, db):
    mk = await group_picker(db, "mg")
    await m.answer("Qaysi guruhdan?", reply_markup=mk) if mk else await m.answer("Guruhlar yo'q.")


@router.callback_query(F.data.startswith("mg:"), STAFF)
async def mv_group(c: CallbackQuery, db):
    studs = await D.group_students(db, int(c.data[3:]))
    b = InlineKeyboardBuilder()
    for s in studs:
        b.button(text=f"{s['last_name']} {s['first_name']} ({s['grade']})", callback_data=f"ms2:{s['id']}")
    b.adjust(1)
    await c.message.answer("Qaysi o'quvchi?" if studs else "Guruhda o'quvchi yo'q.",
                           reply_markup=b.as_markup() if studs else None)
    await c.answer()


@router.callback_query(F.data.startswith("ms2:"), STAFF)
async def mv_student(c: CallbackQuery, db):
    sid = int(c.data[4:])
    s = await db.fetchone("SELECT * FROM students WHERE id=?", (sid,))
    if not s:
        return await c.answer("Topilmadi", show_alert=True)
    groups = await db.fetchall("SELECT * FROM groups WHERE id<>? ORDER BY name", (s["group_id"],))
    if not groups:
        return await c.answer("Boshqa guruh yo'q", show_alert=True)
    b = InlineKeyboardBuilder()
    for g in groups:
        b.button(text=g["name"], callback_data=f"mt:{sid}:{g['id']}")
    b.adjust(2)
    await c.message.answer(f"{s['last_name']} {s['first_name']} qaysi guruhga ko'chirilsin?",
                           reply_markup=b.as_markup())
    await c.answer()


@router.callback_query(F.data.startswith("mt:"), STAFF)
async def mv_do(c: CallbackQuery, db):
    _, sid, gid = c.data.split(":")
    s = await db.fetchone("SELECT * FROM students WHERE id=?", (int(sid),))
    g = await db.fetchone("SELECT * FROM groups WHERE id=?", (int(gid),))
    if not s or not g:
        return await c.answer("Topilmadi", show_alert=True)
    # Joriy oy to'lovlari ham o'quvchi bilan ko'chadi: tushum va oylik yangi guruhga o'tadi,
    # ikkala o'qituvchida bir xil o'quvchi ikki marta hisoblanmaydi. Eski oylar o'zgarmaydi.
    await db.batch([("UPDATE students SET group_id=? WHERE id=?", (g["id"], s["id"])),
                    ("UPDATE payments SET group_id=? WHERE student_id=? AND month=?",
                     (g["id"], s["id"], D.current_month()))])
    await c.message.edit_text(f"✅ {s['last_name']} {s['first_name']} {g['name']} guruhiga ko'chirildi.")
    await c.answer()


# ---------------- Telefon raqamini kiritish / o'zgartirish ----------------
class EditPhone(StatesGroup):
    value = State()


@router.message(F.text == "📞 Telefonni o'zgartirish", ANYONE)
async def ph_start(m: Message, db):
    uid = m.from_user.id
    mk = await group_picker(db, "pg2", uid if await D.get_role(db, uid) == "teacher" else None)
    await m.answer("Qaysi guruh?", reply_markup=mk) if mk else await m.answer("Guruhlar yo'q.")


@router.callback_query(F.data.startswith("pg2:"))
async def ph_group(c: CallbackQuery, db):
    gid = int(c.data[4:])
    if not await can_access_group(db, c.from_user.id, gid):
        return await c.answer("Bu sizning guruhingiz emas", show_alert=True)
    studs = await D.group_students(db, gid)
    b = InlineKeyboardBuilder()
    for s in studs:
        phone = s["phone"] or "(raqam yo'q)"
        b.button(text=f"{s['last_name']} {s['first_name']} {phone}", callback_data=f"ph:{s['id']}")
    b.adjust(1)
    await c.message.answer("Qaysi o'quvchi?" if studs else "Guruhda o'quvchi yo'q.",
                           reply_markup=b.as_markup() if studs else None)
    await c.answer()


@router.callback_query(F.data.startswith("ph:"))
async def ph_student(c: CallbackQuery, db, state: FSMContext):
    sid = int(c.data[3:])
    s = await db.fetchone("SELECT * FROM students WHERE id=?", (sid,))
    if not s or not await can_access_group(db, c.from_user.id, s["group_id"]):
        return await c.answer("Ruxsat yo'q yoki topilmadi", show_alert=True)
    await state.set_state(EditPhone.value)
    await state.update_data(student_id=sid)
    await c.message.answer(f"{s['last_name']} {s['first_name']} uchun yangi telefon raqam (masalan 90 123 45 67):")
    await c.answer()


@router.message(EditPhone.value)
async def ph_save(m: Message, db, state: FSMContext):
    phone = normalize_phone(m.text or "")
    if phone is None:
        return await m.answer("Raqam noto'g'ri. Masalan: 90 123 45 67")
    sid = (await state.get_data())["student_id"]
    await db.execute("UPDATE students SET phone=? WHERE id=?", (phone, sid))
    await state.clear()
    await m.answer(f"✅ Saqlandi: {phone or 'raqamsiz'}")


# ---------------- To'lovni bekor qilish ----------------
@router.message(F.text == "↩️ To'lovni bekor qilish", STAFF)
async def cp_start(m: Message, db):
    mk = await group_picker(db, "cg")
    await m.answer("Qaysi guruh?", reply_markup=mk) if mk else await m.answer("Guruhlar yo'q.")


@router.callback_query(F.data.startswith("cg:"), STAFF)
async def cp_group(c: CallbackQuery, db):
    studs = await D.group_students(db, int(c.data[3:]))
    b = InlineKeyboardBuilder()
    for s in studs:
        b.button(text=f"{s['last_name']} {s['first_name']}", callback_data=f"cs:{s['id']}")
    b.adjust(1)
    await c.message.answer("Qaysi o'quvchi?" if studs else "Guruhda o'quvchi yo'q.",
                           reply_markup=b.as_markup() if studs else None)
    await c.answer()


@router.callback_query(F.data.startswith("cs:"), STAFF)
async def cp_student(c: CallbackQuery, db):
    sid = int(c.data[3:])
    pays = await db.fetchall("SELECT * FROM payments WHERE student_id=? ORDER BY id DESC LIMIT 10", (sid,))
    if not pays:
        return await c.answer("Bu o'quvchida to'lov yo'q", show_alert=True)
    b = InlineKeyboardBuilder()
    for p in pays:
        b.button(text=f"{p['month']} | {money(p['amount'])} | {str(p['created_at'])[:16]}",
                 callback_data=f"cp:{p['id']}")
    b.adjust(1)
    await c.message.answer("Qaysi to'lov bekor qilinsin? (oxirgi 10 ta)", reply_markup=b.as_markup())
    await c.answer()


@router.callback_query(F.data.startswith("cp:"), STAFF)
async def cp_ask(c: CallbackQuery, db):
    p = await db.fetchone("SELECT * FROM payments WHERE id=?", (int(c.data[3:]),))
    if not p:
        return await c.answer("Topilmadi", show_alert=True)
    await c.message.answer(f"⚠️ {p['month']} oyi uchun {money(p['amount'])} to'lov bekor qilinsinmi?",
                           reply_markup=confirm_kb(f"cpy:{p['id']}"))
    await c.answer()


@router.callback_query(F.data.startswith("cpy:"), STAFF)
async def cp_do(c: CallbackQuery, db):
    pid = int(c.data[4:])
    p = await db.fetchone("SELECT * FROM payments WHERE id=?", (pid,))
    if not p:
        return await c.answer("Topilmadi", show_alert=True)
    await db.execute("DELETE FROM payments WHERE id=?", (pid,))
    await c.message.edit_text(f"↩️ {money(p['amount'])} to'lov ({p['month']}) bekor qilindi.")
    await c.answer()


# ---------------- Oylik hisobot (o'tgan oylar) ----------------
@router.message(F.text == "📅 Hisobot", STAFF)
async def rp_start(m: Message):
    await m.answer("Qaysi oy uchun hisobot?", reply_markup=months_markup("rp"))


async def report_text(db, month):
    out = [f"📅 Hisobot: {month}", "", "Guruhlar:"]
    total = 0
    for g in await db.fetchall("SELECT * FROM groups ORDER BY name"):
        studs = await D.group_students(db, g["id"], month)
        paid = sum(1 for s in studs if s["paid"])
        income = await D.group_income(db, g["id"], month)
        total += income
        out.append(f"• {g['name']}: {money(income)} | to'lagan {paid}, qarzdor {len(studs) - paid}")
    out += [f"Jami tushum: {money(total)}", "", "O'qituvchilar oyligi:"]
    paid_out = 0
    for t in await db.fetchall("SELECT u.tg_id, u.full_name FROM users u JOIN teachers t ON t.tg_id=u.tg_id "
                               "ORDER BY u.full_name"):
        sal = (await D.teacher_salary(db, t["tg_id"], month))[3]
        paid_out += sal
        out.append(f"• {t['full_name']}: {money(sal)}")
    out.append(f"Jami oyliklar: {money(paid_out)}")
    return "\n".join(out)


@router.callback_query(F.data.startswith("rp:"), STAFF)
async def rp_month(c: CallbackQuery, db):
    await c.message.edit_text(await report_text(db, c.data[3:]) + "\n\nBoshqa oy:",
                              reply_markup=months_markup("rp"))
    await c.answer()


# ---------------- Oy oxirida qarzdorlar haqida avtomatik xabar ----------------
def local_now():
    return datetime.now(timezone(timedelta(hours=TZ_OFFSET)))


async def send_month_end_notice(bot: Bot, db):
    now = local_now()
    last_day = calendar.monthrange(now.year, now.month)[1]
    key = f"debt-{now:%Y-%m}"
    if now.day != last_day or now.hour < REMINDER_HOUR:
        return
    if await db.fetchone("SELECT 1 FROM notices WHERE key=?", (key,)):
        return
    parts = [f"⏰ Oy tugayapti ({now:%Y-%m}). Qarzdor o'quvchilar:"]
    for g in await db.fetchall("SELECT id FROM groups ORDER BY name"):
        parts.append(await debtors_text(db, g["id"]))
    text = "\n\n".join(parts)
    chunks, cur = [], ""
    for block in text.split("\n\n"):
        if len(cur) + len(block) + 2 > 3800:
            chunks.append(cur)
            cur = ""
        cur += block + "\n\n"
    chunks.append(cur)
    staff = await db.fetchall("SELECT tg_id FROM users WHERE role IN ('admin','director')")
    for u in staff:
        for ch in chunks:
            try:
                await bot.send_message(u["tg_id"], ch.strip())
            except Exception:
                log.warning("Xabar yuborilmadi: %s", u["tg_id"])
    await db.execute("INSERT OR IGNORE INTO notices(key) VALUES(?)", (key,))


async def scheduler(bot: Bot, db):
    while True:
        try:
            await send_month_end_notice(bot, db)
        except Exception:
            log.exception("Oy oxiri xabarida xato")
        await asyncio.sleep(1800)


# ---------------- Qidirish (ism / familiya / telefon) ----------------
class Search(StatesGroup):
    query = State()


@router.message(F.text == "🔎 Qidirish", ANYONE)
async def sr_start(m: Message, state: FSMContext):
    await state.set_state(Search.query)
    await m.answer("Ism, familiya yoki telefon raqamning bir qismini yozing:")


@router.message(Search.query)
async def sr_do(m: Message, db, state: FSMContext):
    q = (m.text or "").strip().casefold()
    if len(q) < 2:
        return await m.answer("Kamida 2 ta belgi yozing.")
    uid = m.from_user.id
    sql = ("SELECT s.*, g.name AS gname FROM students s JOIN groups g ON g.id=s.group_id")
    args = ()
    if await D.get_role(db, uid) == "teacher":
        sql += " WHERE g.teacher_id=?"
        args = (uid,)
    paid_ids = {r["student_id"] for r in await db.fetchall(
        "SELECT DISTINCT student_id FROM payments WHERE month=?", (D.current_month(),))}
    q_digits = "".join(ch for ch in q if ch.isdigit())
    found = []
    for s in await db.fetchall(sql, args):
        names = f"{s['first_name']} {s['last_name']} {s['last_name']} {s['first_name']}".casefold()
        if q in names or (len(q_digits) >= 3 and q_digits in "".join(ch for ch in s["phone"] if ch.isdigit())):
            found.append(s)
    await state.clear()
    if not found:
        return await m.answer("Hech narsa topilmadi.")
    lines = [f"🔎 Topildi: {len(found)} ta" + (" (dastlabki 20 tasi)" if len(found) > 20 else "")]
    for s in found[:20]:
        mark = "✅" if s["id"] in paid_ids else "❌ qarzdor"
        lines.append(f"• {s['last_name']} {s['first_name']} | {s['grade']}-sinf | {s['gname']} | "
                     f"{mark}{with_phone(s)}")
    await m.answer("\n".join(lines))


# ---------------- Excel eksport ----------------
@router.message(F.text == "📥 Excel", STAFF)
async def xl_start(m: Message):
    kb_ = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📋 Qarzdorlar (joriy oy)", callback_data="xl:debt")],
        [InlineKeyboardButton(text="👥 Barcha o'quvchilar", callback_data="xl:stud")],
        [InlineKeyboardButton(text="📅 Oylik hisobot", callback_data="xl:rep")]])
    await m.answer("Qaysi ma'lumotni yuklab olasiz?", reply_markup=kb_)


def _sheet(wb, title, header, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(
            40, max(len(str(c.value or "")) for c in col) + 2)
    return ws


async def _students_rows(db, month, only_debtors):
    rows = []
    for g in await db.fetchall("SELECT * FROM groups ORDER BY name"):
        for s in await D.group_students(db, g["id"], month):
            if only_debtors and s["paid"]:
                continue
            rows.append([g["name"], s["last_name"], s["first_name"], s["grade"], s["phone"],
                         "To'lagan" if s["paid"] else "Qarzdor", s["paid_sum"]])
    return rows


async def _send_xlsx(c: CallbackQuery, wb, name):
    buf = io.BytesIO()
    wb.save(buf)
    await c.message.answer_document(BufferedInputFile(buf.getvalue(), name))
    await c.answer()


STUD_HEAD = ["Guruh", "Familiya", "Ism", "Sinf", "Telefon", "Holat", "To'langan summa"]


@router.callback_query(F.data.in_({"xl:debt", "xl:stud"}), STAFF)
async def xl_students(c: CallbackQuery, db):
    mo = D.current_month()
    debt = c.data == "xl:debt"
    wb = Workbook()
    _sheet(wb, "Qarzdorlar" if debt else "O'quvchilar", STUD_HEAD, await _students_rows(db, mo, debt), True)
    await _send_xlsx(c, wb, f"{'qarzdorlar' if debt else 'oquvchilar'}_{mo}.xlsx")


@router.callback_query(F.data == "xl:rep", STAFF)
async def xl_rep_months(c: CallbackQuery):
    await c.message.answer("Qaysi oy?", reply_markup=months_markup("xr"))
    await c.answer()


@router.callback_query(F.data.startswith("xr:"), STAFF)
async def xl_report(c: CallbackQuery, db):
    mo = c.data[3:]
    wb = Workbook()
    grows = []
    for g in await db.fetchall("SELECT * FROM groups ORDER BY name"):
        studs = await D.group_students(db, g["id"], mo)
        paid = sum(1 for s in studs if s["paid"])
        grows.append([g["name"], await D.group_income(db, g["id"], mo), paid, len(studs) - paid])
    _sheet(wb, "Guruhlar", ["Guruh", "Tushum", "To'laganlar", "Qarzdorlar"], grows, True)
    trows = []
    for t in await db.fetchall("SELECT u.tg_id, u.full_name FROM users u JOIN teachers t ON t.tg_id=u.tg_id "
                               "ORDER BY u.full_name"):
        pt, pv, _, total = await D.teacher_salary(db, t["tg_id"], mo)
        trows.append([t["full_name"], "Foiz" if pt == "percent" else "O'quvchi boshiga", pv, round(total)])
    _sheet(wb, "Oyliklar", ["O'qituvchi", "Turi", "Qiymat", "Oylik"], trows)
    _sheet(wb, "O'quvchilar", STUD_HEAD, await _students_rows(db, mo, False))
    await _send_xlsx(c, wb, f"hisobot_{mo}.xlsx")


# ---------------- Direktor statistikasi ----------------
def _pct(a, b):
    return f"{a * 100 / b:.0f}%" if b else "-"


@router.message(F.text == "📈 Statistika", DIRECTOR)
async def stats(m: Message, db):
    cur, prev = prev_months(2)
    income = {mo: (await db.fetchone("SELECT COALESCE(SUM(amount),0) t FROM payments WHERE month=?", (mo,)))["t"]
              for mo in (cur, prev)}
    n_students = (await db.fetchone("SELECT COUNT(*) n FROM students"))["n"]
    n_paid = (await db.fetchone(
        "SELECT COUNT(DISTINCT p.student_id) n FROM payments p JOIN students s ON s.id=p.student_id "
        "WHERE p.month=?", (cur,)))["n"]
    salaries = 0
    for t in await db.fetchall("SELECT tg_id FROM teachers"):
        salaries += (await D.teacher_salary(db, t["tg_id"], cur))[3]
    diff = income[cur] - income[prev]
    out = [f"📈 Statistika ({cur})", "",
           f"💵 Tushum: {money(income[cur])}",
           f"   O'tgan oy ({prev}): {money(income[prev])} ({'+' if diff >= 0 else '-'}{money(abs(diff))})",
           f"👨‍🏫 O'qituvchilar oyligi: {money(salaries)}",
           f"🏦 Qolgan (tushum - oylik): {money(income[cur] - salaries)}", "",
           f"👥 O'quvchilar: {n_students} ta",
           f"✅ To'laganlar: {n_paid} ({_pct(n_paid, n_students)})",
           f"❌ Qarzdorlar: {n_students - n_paid} ({_pct(n_students - n_paid, n_students)})", ""]
    att = await db.fetchone("SELECT COALESCE(SUM(present),0) p, COUNT(*) n FROM attendance WHERE day LIKE ?",
                            (cur + "-%",))
    out.append(f"📋 Davomat (shu oy): kelish {_pct(att['p'], att['n'])}" if att["n"]
               else "📋 Davomat: shu oy hali olinmagan")
    out += ["", "Guruhlar:"]
    for g in await db.fetchall("SELECT * FROM groups ORDER BY name"):
        studs = await D.group_students(db, g["id"], cur)
        debt = sum(1 for s in studs if not s["paid"])
        ga = await db.fetchone("SELECT COALESCE(SUM(present),0) p, COUNT(*) n FROM attendance "
                               "WHERE group_id=? AND day LIKE ?", (g["id"], cur + "-%"))
        out.append(f"• {g['name']}: {money(await D.group_income(db, g['id'], cur))} | "
                   f"o'quvchi {len(studs)}, qarzdor {debt} | davomat {_pct(ga['p'], ga['n'])}")
    await m.answer("\n".join(out))


# ---------------- Adminni o'chirish (faqat direktor) ----------------
@router.message(F.text == "🗑 Adminni o'chirish", DIRECTOR)
async def xa_start(m: Message, db):
    admins = await db.fetchall("SELECT tg_id, full_name FROM users WHERE role='admin' ORDER BY full_name")
    if not admins:
        return await m.answer("Adminlar yo'q.")
    b = InlineKeyboardBuilder()
    for a in admins:
        b.button(text=f"{a['full_name']} ({a['tg_id']})", callback_data=f"xa:{a['tg_id']}")
    b.adjust(1)
    await m.answer("Qaysi admin o'chirilsin?", reply_markup=b.as_markup())


@router.callback_query(F.data.startswith("xa:"), DIRECTOR)
async def xa_ask(c: CallbackQuery, db):
    a = await db.fetchone("SELECT * FROM users WHERE tg_id=? AND role='admin'", (int(c.data[3:]),))
    if not a:
        return await c.answer("Topilmadi", show_alert=True)
    await c.message.answer(f"⚠️ Admin {a['full_name']} o'chirilsinmi? Botdan foydalanish huquqi olinadi.",
                           reply_markup=confirm_kb(f"xay:{a['tg_id']}"))
    await c.answer()


@router.callback_query(F.data.startswith("xay:"), DIRECTOR)
async def xa_do(c: CallbackQuery, db):
    a = await db.fetchone("SELECT * FROM users WHERE tg_id=? AND role='admin'", (int(c.data[4:]),))
    if not a:
        return await c.answer("Topilmadi", show_alert=True)
    await db.execute("DELETE FROM users WHERE tg_id=? AND role='admin'", (a["tg_id"],))
    await c.message.edit_text(f"🗑 Admin {a['full_name']} o'chirildi.")
    await c.answer()


# ---------------- Guruhlar (o'qituvchi -> guruhlar -> o'quvchilar) ----------------
TEACHER_ONLY = Role("teacher")


async def _groups_markup(db, teacher_id, back=False):
    """teacher_id=0 - o'qituvchisiz guruhlar."""
    if teacher_id:
        groups = await db.fetchall("SELECT * FROM groups WHERE teacher_id=? ORDER BY name", (teacher_id,))
    else:
        groups = await db.fetchall("SELECT * FROM groups WHERE teacher_id IS NULL ORDER BY name")
    counts = {r["group_id"]: r["n"] for r in await db.fetchall(
        "SELECT group_id, COUNT(*) n FROM students GROUP BY group_id")}
    b = InlineKeyboardBuilder()
    for g in groups:
        b.button(text=f"{g['name']} ({counts.get(g['id'], 0)} ta)", callback_data=f"gl:g:{g['id']}")
    b.adjust(2)
    if back:
        b.row(InlineKeyboardButton(text="⬅️ O'qituvchilar", callback_data="gl:back"))
    return (b.as_markup() if groups else None), bool(groups)


@router.message(F.text == "📚 Guruhlarim", TEACHER_ONLY)
async def my_groups(m: Message, db):
    mk, ok = await _groups_markup(db, m.from_user.id)
    await m.answer("Guruhingizni tanlang:", reply_markup=mk) if ok else await m.answer("Sizda guruh yo'q.")


async def _teachers_markup(db):
    teachers = await db.fetchall(
        "SELECT u.tg_id, u.full_name, (SELECT COUNT(*) FROM groups g WHERE g.teacher_id=u.tg_id) n "
        "FROM users u JOIN teachers t ON t.tg_id=u.tg_id ORDER BY u.full_name")
    b = InlineKeyboardBuilder()
    for t in teachers:
        b.button(text=f"{t['full_name']} ({t['n']})", callback_data=f"gl:t:{t['tg_id']}")
    b.adjust(1)
    orphan = (await db.fetchone("SELECT COUNT(*) n FROM groups WHERE teacher_id IS NULL"))["n"]
    if orphan:
        b.button(text=f"⚠️ O'qituvchisiz guruhlar ({orphan})", callback_data="gl:t:0")
    b.adjust(1)
    return b.as_markup() if (teachers or orphan) else None


@router.message(F.text == "📚 Guruhlar", STAFF)
async def all_groups(m: Message, db):
    mk = await _teachers_markup(db)
    await m.answer("O'qituvchini tanlang:", reply_markup=mk) if mk else await m.answer("Guruhlar yo'q.")


@router.callback_query(F.data == "gl:back", STAFF)
async def gl_back(c: CallbackQuery, db):
    mk = await _teachers_markup(db)
    await c.message.edit_text("O'qituvchini tanlang:", reply_markup=mk)
    await c.answer()


@router.callback_query(F.data.startswith("gl:t:"), STAFF)
async def gl_teacher(c: CallbackQuery, db):
    tid = int(c.data[5:])
    mk, ok = await _groups_markup(db, tid, back=True)
    if tid:
        t = await db.fetchone("SELECT full_name FROM users WHERE tg_id=?", (tid,))
        title = f"👨‍🏫 {t['full_name']} guruhlari:" if t else "Guruhlar:"
    else:
        title = "O'qituvchisiz guruhlar:"
    await c.message.edit_text(title if ok else title + "\nGuruh yo'q.", reply_markup=mk or InlineKeyboardMarkup(
        inline_keyboard=[[InlineKeyboardButton(text="⬅️ O'qituvchilar", callback_data="gl:back")]]))
    await c.answer()


@router.callback_query(F.data.startswith("gl:g:"))
async def gl_group(c: CallbackQuery, db):
    gid = int(c.data[5:])
    if not await can_access_group(db, c.from_user.id, gid):
        return await c.answer("Bu sizning guruhingiz emas", show_alert=True)
    g = await db.fetchone("SELECT g.name, u.full_name tname FROM groups g "
                          "LEFT JOIN users u ON u.tg_id=g.teacher_id WHERE g.id=?", (gid,))
    studs = await D.group_students(db, gid)
    paid = sum(1 for s in studs if s["paid"])
    tname = g["tname"] or "o'qituvchisiz"
    head = [f"📚 {g['name']} | {tname}",
            f"O'quvchilar: {len(studs)} | to'lagan: {paid} | qarzdor: {len(studs) - paid}", ""]
    lines = [f"{i}. {'✅' if s['paid'] else '❌'} {s['last_name']} {s['first_name']} "
             f"({s['grade']}-sinf){with_phone(s)}" for i, s in enumerate(studs, 1)]
    text = "\n".join(head + (lines or ["Guruhda o'quvchi yo'q."]))
    for i in range(0, len(text), 4000):
        await c.message.answer(text[i:i + 4000])
    await c.answer()
